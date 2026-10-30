"""AidData（William & Mary）China's Global Loans and Grants Dataset, Version 1.0（CLG-Global 1.0，2025-11-18 發布）。

- 內容：2000–2023 年中國官方機構（政府、政策性銀行、國有企業…）對 217 個低、中、高所得國家與地區的開發融資，
  一筆一個項目（贈款、貸款、債務重整…），33,580 筆；執行細節更新到 2025 年。由中低所得國的 CLG-LMIC 1.0 與
  高所得國的 CLG-HIC 1.0 合併而成（每筆的 Source_Dataset 欄）。取代原本的 GCDF 3.0（2000–2021，165 個中低所得國）。
- 授權：Open Data Commons Attribution License（ODC-By 1.0）。xlsx 的 License 工作表：「The original dataset produced by
  AidData is licensed using the Open Data Commons Attribution License (ODC-By). Any use of the dataset and its content
  must include appropriate attribution to AidData.」「you must make clear to others the license of the database and keep
  intact any notices on the original database」（DESIGN.md 第 8 節第 12、91、209 點）。
- 引用（xlsx 的 Contents 工作表，非學術用途）：AidData. 2025. China's Global Loans and Grants Dataset, Version 1.0.
  Retrieved from www.aiddata.org/data/chinas-global-loans-and-grants-dataset-1-0.
- 下載：zip（43.8 MB，不必填表、不必登入），裡面的 xlsx 40.6 MB、欄位定義與 TUFF 4.0 方法論的 PDF；`__MACOSX/` 略過。
  openpyxl 的 read_only 模式讀 Records 工作表約 1 分鐘，讀完只留用到的欄存成 .pkl 快取。
"""
from __future__ import annotations

import io
import pickle
import zipfile
from functools import lru_cache

from common import CACHE_DIR, download

URL = "https://docs.aiddata.org/ad4/datasets/AidDatas_CLG_Global_Dataset_v1.0.zip"
ZIP = "clg1.zip"
PAGE = "https://www.aiddata.org/data/chinas-global-loans-and-grants-dataset-1-0"
SHEET = "CLG-Global 1.0_Records"

# 用到的欄位（xlsx 的欄名）
COLUMNS = [
    "AidData_Record_ID",
    "Recommended_for_Aggregates",
    "Country_of_Activity",
    "Country_of_Activity_ISO3",
    "Commitment_Year",
    "Title",
    "Status",
    "Flow_Type_Simplified",
    "Flow_Class",
    "Sector_Code",
    "Adjusted_Amount_Constant_USD_2023",
    "Adjusted_Amount_Nominal_USD",
    "Source_Dataset",
]


def _workbook():
    import openpyxl

    path = download(URL, "aiddata", ZIP)
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.endswith(".xlsx") and not n.startswith("__MACOSX"))
        data = z.read(name)
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


@lru_cache(maxsize=1)
def _load() -> dict:
    """{"records": Records 工作表的每一筆（只留 COLUMNS，欄名當 key）, "countries": CountryList 的 (英文名稱, 筆數)}"""
    cache = CACHE_DIR / "aiddata" / "clg1-records.pkl"
    if cache.exists():
        with open(cache, "rb") as f:
            return pickle.load(f)
    wb = _workbook()
    rows = wb[SHEET].iter_rows(values_only=True)
    header = next(rows)
    index = {name: header.index(name) for name in COLUMNS}
    data = {
        "records": [{name: row[i] for name, i in index.items()} for row in rows if row[0] is not None],
        # 欄位：Country、OECD Income Status、ODA eligibility、Record Count；筆數欄沒有項目的寫「No projects/activities found」，
        # 表格最後兩列是說明文字（其他欄空白）
        "countries": [
            (row[0].strip(), row[3])
            for row in wb["CountryList"].iter_rows(values_only=True)
            if row and row[0] and len(row) > 3 and row[3] is not None and row[0] != "Country"
        ],
    }
    with open(cache, "wb") as f:
        pickle.dump(data, f)
    return data


def records() -> list[dict]:
    """Records 工作表的每一筆"""
    return _load()["records"]


def covered_countries() -> list[tuple[str, int | str]]:
    """CountryList 工作表：AidData 有系統地搜尋過的 217 個國家與地區，(英文名稱, 筆數或「No projects/activities found」)"""
    return _load()["countries"]
