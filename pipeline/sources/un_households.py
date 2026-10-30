"""UN DESA Population Division 的家戶資料庫（CC BY 3.0 IGO）：
Database on Household Size and Composition（2022、2026 版）與 Database on the Households and Living Arrangements
of Older Persons 2026。

- 都是一個 Excel 檔，資料工作表一列是一個國家的一次普查或調查：國名、Location ID、（2026 版）ISO3 代碼……
  資料來源類別（DYB、IPUMS、LFS、DHS、MICS）、參考日期（年份）、各項比例。
- 2022 版的參考日期有的是日期格式，有的是「19/10/2015」這種字串；第二欄是 ISO 數字代碼，科索沃用 412。
- 2026 版的數字大多存成字串，沒有值寫「..」；有幾列的來源類別是「DYB & China Population Census Yearbook 2020」
  「IPUMS & Russian Federation 2021 Census」這種合寫，或各國普查的全名（埃及、摩洛哥、巴布亞紐幾內亞……），
  依第一個詞判斷，各國普查跟 DYB 一樣算普查。
- 同一國同一年可能有好幾個來源（例如普查同時在 DYB 與 IPUMS），依 SOURCE_PRIORITY 取一個：
  普查（DYB、IPUMS）優先，其次是歐盟勞動力調查（LFS），最後才是抽樣調查（DHS、MICS）。
- 兩個資料庫都沒有台灣。
- 住宅與房價主題的平均每戶人數用 2026 版（2026-10-30 起，原本是 2022 版）；婚姻與家庭主題用 2026 版的單人家戶、單親、三代同堂、隔代家戶的比例，
  以及老人資料庫裡 60 歲以上的人獨居（「1 member」）的比例。
- 三代同堂、隔代家戶要有戶內成員關係的個體資料（IPUMS、DHS、MICS）才算得出來，DYB 的表沒有，
  所以只有普查表的國家（日本、韓國、德國……）沒有這兩項。
"""
from __future__ import annotations

import datetime

import openpyxl

from common import download
from sources.countries import from_numeric

FILES = "https://www.un.org/development/desa/pd/sites/www.un.org.development.desa.pd/files/"
URL = FILES + "undesa_pd_2022_hh-size-composition.xlsx"
SHEET = "HH size and composition 2022"
URL_2026 = FILES + "undesa_pd_2026_hh-size-composition.xlsx"
SHEET_2026 = "HH size and composition 2026"
OLDER_URL_2026 = FILES + "undesa_pd_2026_living-arrangements-older-persons.xlsx"
OLDER_SHEET_2026 = "HHLA of Older Persons 2026"
SOURCE_PRIORITY = ["DYB", "IPUMS", "LFS", "DHS", "MICS"]
# 2026 版家戶檔的欄位（標題列）
COMPOSITION_2026 = {
    "one-person": "1 member",  # 「One-person」那一欄有幾國是空的（中國 2020 年），「1 member」相同而且較齊全
    "single-parent": "Single parent with children",
    "three-generation": "Three generation",
    "skip-generation": "Skip generation",
}

Series = dict[str, dict[int, float]]


def _year(cell: object) -> int | None:
    if isinstance(cell, (datetime.date, datetime.datetime)):
        return cell.year
    text = str(cell).strip()
    return int(text[-4:]) if text[-4:].isdigit() else None


def _number(cell: object) -> float | None:
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell).strip())
    except ValueError:  # 「..」或空白
        return None


def _rank(source: object) -> int:
    first = str(source or "").split(" ")[0]
    if first in SOURCE_PRIORITY:
        return SOURCE_PRIORITY.index(first)
    return 0 if "census" in str(source).lower() or "recensement" in str(source).lower() else len(SOURCE_PRIORITY)


def _pick(rows: list[tuple[str, int, int, float]]) -> Series:
    """(iso3, 年, 來源的順位, 值) → 每國每年取順位最前面的一個"""
    best: dict[tuple[str, int], tuple[int, float]] = {}
    for iso3, year, rank, value in rows:
        if (iso3, year) not in best or rank < best[(iso3, year)][0]:
            best[(iso3, year)] = (rank, value)
    out: Series = {}
    for (iso3, year), (_, value) in best.items():
        out.setdefault(iso3, {})[year] = value
    return out


def average_household_size() -> Series:
    """{iso3: {year: 平均每戶人數}}（2026 版；製作手冊第八批 150 以前是 2022 版的 _column(4, …)）"""
    names, rows = _rows_2026(URL_2026, "hh-size-composition-2026.xlsx", SHEET_2026)
    iso, src, ref = names.index("ISO3 Code"), names.index("Data source category"), names.index("Reference year")
    col = names.index("Average household size (number of members)")
    picked = []
    for row in rows:
        value, year = _number(row[col]), _year(row[ref])
        if value is not None and year is not None and row[iso]:
            picked.append((str(row[iso]), year, _rank(row[src]), value))
    return _pick(picked)


def one_person_share() -> Series:
    """{iso3: {year: 單人家戶占所有家戶的比例（%）}}，2022 版「Households by size」的「1 member」。"""
    return _column(5, "1 member")


def _column(index: int, title: str) -> Series:
    path = download(URL, "undesa", "hh-size-composition-2022.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Country or area")
    if not str(rows[header][index]).startswith(title):
        raise RuntimeError(f"{SHEET} 的欄位變了：{rows[header][:6]}")
    picked = []
    for row in rows[header + 1 :]:
        code, source, date, size = row[1], row[2], row[3], row[index]
        if code is None or not isinstance(size, (int, float)):
            continue
        iso3 = "XKX" if int(code) == 412 else from_numeric(code)
        year = _year(date)
        if iso3 and year is not None:
            picked.append((iso3, year, _rank(source), float(size)))
    return _pick(picked)


def _rows_2026(url: str, filename: str, sheet: str) -> tuple[list[str], list[tuple]]:
    path = download(url, "undesa", filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Country and area")
    names = [str(c).strip() if c is not None else "" for c in rows[header]]
    for need in ("ISO3 Code", "Data source category", "Reference year"):
        if need not in names:
            raise RuntimeError(f"{sheet} 的欄位變了，找不到「{need}」")
    return names, [r for r in rows[header + 1 :] if r and r[0]]


def composition_2026() -> dict[str, Series]:
    """2026 版的家戶組成：{指標: {iso3: {year: 占所有家戶的比例（%）}}}，指標見 COMPOSITION_2026。"""
    names, rows = _rows_2026(URL_2026, "hh-size-composition-2026.xlsx", SHEET_2026)
    iso, src, ref = names.index("ISO3 Code"), names.index("Data source category"), names.index("Reference year")
    out = {}
    for key, title in COMPOSITION_2026.items():
        col = names.index(title)
        picked = []
        for row in rows:
            value, year = _number(row[col]), _year(row[ref])
            if value is not None and year is not None and row[iso]:
                picked.append((str(row[iso]), year, _rank(row[src]), value))
        out[key] = _pick(picked)
    return out


def older_alone_2026() -> Series:
    """60 歲以上的人獨居（戶內只有 1 人）的比例（%，男女合計）：{iso3: {year: 比例}}"""
    names, rows = _rows_2026(OLDER_URL_2026, "living-arrangements-older-persons-2026.xlsx", OLDER_SHEET_2026)
    iso, src, ref = names.index("ISO3 Code"), names.index("Data source category"), names.index("Reference year")
    age, sex, alone = names.index("Age range"), names.index("Sex"), names.index("1 member")
    picked = []
    for row in rows:
        if row[age] != "60 or over" or row[sex] != "Both sexes":
            continue
        value, year = _number(row[alone]), _year(row[ref])
        if value is not None and year is not None and row[iso]:
            picked.append((str(row[iso]), year, _rank(row[src]), value))
    return _pick(picked)
