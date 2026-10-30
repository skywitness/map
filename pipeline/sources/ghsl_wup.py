"""JRC GHSL 的 GHS-WUP R2025A（CC BY 4.0，© European Union）：給聯合國《World Urbanization Prospects 2025》做的
城市與各國城鄉人口，1950–2100 每 5 年一期（31 期），取代原本的 GHS-UCDB R2024A 與 GHS-COUNTRY-STATS R2024A。

- 方法：GHS-WUP 是 GHSL 2023 資料包（P2023）之外另一套「不相容的資料流」：1975–2020 年的人口網格重新評估、配合聯合國
  WPP 2024 的各國人口，2025–2100 年用 CRISP 模型推估建物與人口網格，每一期套用「城鄉程度」（Degree of Urbanisation）
  劃出城市（urban centre：每 km² 1,500 人以上的 1 km 格子連成一片、合計 5 萬人以上）。1950–1970 年沒有網格，
  是用 WUP 的資料回推（線性迴歸或各國城市人口的比例，`WUP` 欄），所以那幾期沒有面積與建物面積。
- `MTUC`：GHS-WUP-MTUC R2025A V1.1（2026-02-13 修正 752 個城市的名稱）的 `UC_STATS` 工作表，一個城市一期一列：
  `ID_UC_G0`（同一個城市每一期都一樣）、`UCname`（主要名稱，GISCO、OSM 或人工）、`UNLocID`（聯合國的國家代碼，
  M49；香港、澳門、台灣分開）、`AREA_km2`、`POP`、`BU_km2`、人口加權的中心點（`Lat`、`Lon`，WGS84）、`CapitalFlag`。
  城市在 2025 年以前「消失」（併進別的城市）或之後才「出現」的，那幾期沒有列。
- `STATS`：GHS-WUP-COUNTRY-STATS R2025A V1.0 的 `Country_STATS_L1` 工作表：每個國家（UNLocID）、每一期住在城市（UC）、
  城鎮與半密集區（UCL）、鄉村（RUR）的人口、面積與建物面積（面積與建物面積只有 1975 年起）。
- 引用：GHSL 的「Use conditions and how to cite」要寫最新的 GHSL 參考文獻；下面兩個資料集的引用照 xlsx 的 Description。

xlsx 用 openpyxl 的 read_only 模式讀（MTUC 36 萬列約 1 分鐘），讀完存成 .pkl 快取。
"""
from __future__ import annotations

import io
import pickle
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache

from common import CACHE_DIR, download
from sources.countries import from_numeric

BASE = "https://jeodpp.jrc.ec.europa.eu/ftp/jrc-opendata/GHSL"
MTUC = f"{BASE}/GHS_WUP_MTUC_GLOBE_R2025A/V1-1/GHS_WUP_MTUC_GLOBE_R2025A_V1_1_statistics.zip"
MTUC_XLSX = "GHS_WUP_MTUC_MT_GLOBE_R2025A_v1_1.xlsx"
STATS = f"{BASE}/GHS_WUP_COUNTRY_STATS_GLOBE_R2025A/V1-0/GHS_WUP_COUNTRY_STATS_GLOBE_R2025A_V1_0.zip"
STATS_XLSX = "GHS_WUP_COUNTRY_STATS_MT_GLOBE_R2025A_V1_0.xlsx"
CACHE = "ghsl-wup"
PAGE = "https://human-settlement.emergency.copernicus.eu/ghs_wup2025.php"

EPOCHS = list(range(1950, 2101, 5))
GRID_FROM = 1975  # 1975 年起才有網格（面積、建物面積）；之前是回推
PROJECTION_FROM = 2025  # 2025 年起是推估

CITATION_MTUC = (
    "Schiavina, M., Alessandrini, A., Melchiorri, M., Dijkstra, L. (2025): GHS-WUP-MTUC R2025A - GHS-WUP multitemporal "
    "urban centres, obtained from the Degree of Urbanisation grids (GHS-WUP-DEGURBA R2025A) and linked across epochs, "
    "multitemporal (1950-2100). European Commission, Joint Research Centre (JRC). V1.1 (2026-02-13)."
)
CITATION_STATS = (
    "Schiavina, M., Alessandrini, A., Melchiorri, M., Jacobs-Crisioni, C., Dijkstra, L. (2025): GHS-WUP-COUNTRY-STATS "
    "R2025A - GHS-WUP country statistics by Degree of Urbanisation classes, multitemporal (1950-2100). European "
    "Commission, Joint Research Centre (JRC). doi:10.2905/ee52f334-aaee-4b9e-aa00-d38ecf6c11a9"
)
CITATION_PACKAGE = (
    "Schiavina, M., Melchiorri, M., Mari Rivero, I., Florio, P., Freire, S. et al. (2025): GHSL WUP Projections Data "
    "Package 2025 - Public release GHS-WUP R2025. Publications Office of the European Union, Luxembourg. "
    "doi:10.2760/2416436, JRC144209"
)


def to_iso3(loc_id: int) -> str | None:
    """聯合國的 LocID（M49）→ ISO alpha-3；海峽群島（830）等沒有 ISO 代碼的回傳 None"""
    return from_numeric(loc_id)


@dataclass
class UrbanCentre:
    id: int  # ID_UC_G0
    name: str
    loc_id: int  # 2025 年（或最接近的一期）所在國家的 UNLocID
    capital: bool
    lon: float
    lat: float
    pop: dict[int, float] = field(default_factory=dict)  # 期 → 人口（那一期的城市範圍）；還不是城市的期沒有
    area: dict[int, float] = field(default_factory=dict)  # 期 → km²（1975 年起）


def _xlsx(url: str, member: str):
    import openpyxl

    path = download(url, CACHE, url.rsplit("/", 1)[1])
    with zipfile.ZipFile(path) as z:
        data = z.read(member)
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


@lru_cache(maxsize=1)
def urban_centres() -> list[UrbanCentre]:
    """每個城市（ID_UC_G0）1950–2100 每一期的人口與面積；中心點、名稱、國家用最接近 2025 年的那一期。"""
    cache = CACHE_DIR / CACHE / "mtuc-v1_1.pkl"
    if cache.exists():
        with open(cache, "rb") as f:
            return pickle.load(f)
    rows = _xlsx(MTUC, MTUC_XLSX)["UC_STATS"].iter_rows(values_only=True)
    header = next(rows)
    col = {name: header.index(name) for name in header}
    centres: dict[int, UrbanCentre] = {}
    best: dict[int, int] = {}  # 城市 → 用來取名稱與中心點的那一期跟 2025 年差幾年
    for r in rows:
        if r[0] is None:
            continue
        cid, year, pop = int(r[col["ID_UC_G0"]]), int(r[col["Year"]]), r[col["POP"]]
        c = centres.get(cid)
        if c is None:
            c = centres[cid] = UrbanCentre(cid, "", 0, False, 0.0, 0.0)
        if pop and pop > 0:
            c.pop[year] = float(pop)
        area = r[col["AREA_km2"]]
        if area and area > 0:
            c.area[year] = float(area)
        gap = abs(year - PROJECTION_FROM)
        if cid not in best or gap < best[cid]:
            best[cid] = gap
            c.name = (r[col["UCname"]] or "").strip()
            c.loc_id = int(r[col["UNLocID"]])
            c.capital = bool(r[col["CapitalFlag"]])
            c.lat, c.lon = float(r[col["Lat"]]), float(r[col["Lon"]])
    out = [c for c in centres.values() if c.pop]
    cache.parent.mkdir(parents=True, exist_ok=True)
    with open(cache, "wb") as f:
        pickle.dump(out, f)
    return out


@dataclass
class CountryStats:
    names: dict[int, str]  # UNLocID → 聯合國的英文國名
    pop: dict[int, dict[int, tuple[float, float, float]]]  # UNLocID → 期 → (城市, 城鎮與半密集區, 鄉村)
    built: dict[int, dict[int, float]]  # UNLocID → 期 → 建物面積（km²，城鄉加總；1975 年起）


@lru_cache(maxsize=1)
def country_stats() -> CountryStats:
    rows = _xlsx(STATS, STATS_XLSX)["Country_STATS_L1"].iter_rows(values_only=True)
    header = next(rows)
    col = {name: header.index(name) for name in header}
    names: dict[int, str] = {}
    parts: dict[int, dict[int, dict[str, float]]] = {}
    built: dict[int, dict[int, float]] = {}
    for r in rows:
        if r[0] is None:
            continue
        loc, year, cls = int(r[col["UNLocID"]]), int(r[col["Year"]]), str(r[col["DEGURBA"]])
        names[loc] = str(r[col["UNLocName"]])
        parts.setdefault(loc, {}).setdefault(year, {})[cls] = float(r[col["POP"]] or 0)
        bu = r[col["BU_km2"]]
        if bu is not None:
            built.setdefault(loc, {})[year] = built.get(loc, {}).get(year, 0.0) + float(bu)
    pop = {}
    for loc, by_year in parts.items():
        years = sorted(by_year)
        if years != EPOCHS:
            raise RuntimeError(f"GHS-WUP-COUNTRY-STATS 的 {names[loc]} 年份不是 1950–2100 每 5 年：{years[:3]}…")
        pop[loc] = {y: (v.get("UC", 0.0), v.get("UCL", 0.0), v.get("RUR", 0.0)) for y, v in by_year.items()}
    return CountryStats(names=names, pop=pop, built=built)


VECTOR = f"{BASE}/GHS_WUP_MTUC_GLOBE_R2025A/V1-1/GHS_WUP_MTUC_GLOBE_R2025A_V1_1_vector.zip"
VECTOR_SHP = "GHS_WUP_MTUC_MT_GLOBE_R2025A_v1_1"


def _dbf(data: bytes) -> list[dict[str, str]]:
    """dBase III 的欄位（只有文字與數字欄，這裡全部當字串讀）"""
    import struct

    n, header_len, record_len = struct.unpack("<IHH", data[4:12])
    fields, pos = [], 32
    while data[pos] != 0x0D:
        fields.append((data[pos:pos + 11].split(b"\0")[0].decode(), data[pos + 16]))
        pos += 32
    out = []
    for i in range(n):
        rec = data[header_len + i * record_len: header_len + (i + 1) * record_len]
        row, at = {}, 1  # 第一個位元組是刪除標記
        for name, length in fields:
            row[name] = rec[at:at + length].decode("latin1").strip()
            at += length
        out.append(row)
    return out


def boundaries():
    """GHS-WUP-MTUC 的城市範圍（vector.zip 的 shapefile，Mollweide ESRI:54009，每個城市每一期一個多邊形，1975–2100）：
    依序產生 (ID_UC_G0, 期, [環 (N×2 的 numpy 陣列，公尺)])。標準函式庫讀 shp（多邊形的格式很單純），不必裝 GDAL。"""
    import struct

    import numpy as np

    path = download(VECTOR, CACHE, VECTOR.rsplit("/", 1)[1])
    with zipfile.ZipFile(path) as z:
        records = _dbf(z.read(f"{VECTOR_SHP}.dbf"))
        shp = z.read(f"{VECTOR_SHP}.shp")
    pos = 100
    for rec in records:
        _num, length = struct.unpack(">ii", shp[pos:pos + 8])
        body = shp[pos + 8: pos + 8 + length * 2]
        pos += 8 + length * 2
        if struct.unpack("<i", body[:4])[0] == 0:
            continue  # 空的形狀
        nparts, npts = struct.unpack("<ii", body[36:44])
        parts = list(struct.unpack(f"<{nparts}i", body[44:44 + 4 * nparts])) + [npts]
        start = 44 + 4 * nparts
        pts = np.frombuffer(body[start:start + 16 * npts], dtype="<f8").reshape(-1, 2)
        yield int(float(rec["ID_UC_G0"])), int(float(rec["Year"])), [pts[parts[k]:parts[k + 1]] for k in range(nparts)]
