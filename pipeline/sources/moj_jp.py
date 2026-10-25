"""日本法務省《令和 7 年版 犯罪白書》2-4-2-1 圖「刑事施設の年末収容人員・人口比の推移」的 Excel：
日本的在監人數，只用在治安與司法主題的監禁人口率。

- 刑事施設（刑務所、少年刑務所、拘置所）每年 12 月 31 日的收容人員「総数_計」，1946–2024 年，
  含受刑者、死刑確定者、被告人（未決拘禁）、被疑者、労役場留置者與其他（引致状による留置者等）。
  資料來自行刑統計年報、矯正統計年報（e-Stat 的矯正統計）。
- 不含少年院、少年鑑別所（少年的保護處分）與入管的收容設施。
- 白書自己的「年末人口比」用總務省的人口；這裡改用 WPP 的年中人口。

授權：法務省網站的內容適用「公共データ利用規約（第 1.0 版）」（PDL1.0，https://www.moj.go.jp/hisho06_00280.html），
跟 CC BY 4.0 相容，要註明出處，編輯、加工過的要另外寫明。
"""
from __future__ import annotations

import re

import openpyxl

from common import download

EDITION = 72  # 令和 7 年版
URL = f"https://hakusyo1.moj.go.jp/jp/{EDITION}/nfm/excel/2-4-2-1.xlsx"
PAGE = f"https://hakusyo1.moj.go.jp/jp/{EDITION}/nfm/n{EDITION}_2_2_4_2_1.html"


def _number(cell: object) -> float | None:
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell or "").replace(",", "").strip()
    return float(text) if re.fullmatch(r"\d+(\.\d+)?", text) else None


def year_end_inmates() -> dict[int, float]:
    """{年: 刑事施設的年末收容人員}"""
    path = download(URL, "moj-jp", f"hakusyo{EDITION}-2-4-2-1.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and "年次（西暦）" in [str(c) for c in r])
    cols = [str(c) if c is not None else "" for c in rows[head]]
    year_col, total_col = cols.index("年次（西暦）"), cols.index("総数_計")
    out: dict[int, float] = {}
    for r in rows[head + 1:]:
        m = re.fullmatch(r"(\d{4})年", str(r[year_col] or "").strip())
        total = _number(r[total_col])
        if m and total is not None:
            out[int(m.group(1))] = total
    if len(out) < 50:
        raise RuntimeError(f"犯罪白書 2-4-2-1 只讀到 {len(out)} 年")
    return out
