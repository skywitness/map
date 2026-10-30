"""WHO 的道路安全資料：死亡率（GHO）、各類道路使用者的死亡占比與登記車輛數（Global Status Report on Road Safety）。

- 死亡率：GHO 的 `RS_198`（Estimated road traffic death rate per 100 000 population，兩性合計），2000–2023 年，
  WHO 的估計值（30 天定義：事故後 30 天內死亡），所有會員國都有、沒有台灣。GHO 是 CC BY 4.0（data.who.int 的條款）。
- 各類道路使用者的死亡占比（各國通報的死亡人數，依道路使用者分）：
  - 2016 年前後：GHO 的 `RS_246`（Global Status Report on Road Safety 2018 的調查），CC BY 4.0。
  - 2021 年前後：Global Status Report on Road Safety 2023 的指標檔（GSRRS 2023，`CP_18a`–`CP_18e`），
    檔案的 Notes 寫明引用《Global status report on road safety 2023: country and territory profiles》，
    授權 CC BY-NC-SA 3.0 IGO（WHO 出版品的授權）。
- 登記的機動車輛數：2016 年前後用 GHO 的 `RS_194`（CC BY 4.0），2021 年前後用 GSRRS 2023 的 `CP_36a`（CC BY-NC-SA 3.0 IGO）；
  其中的二、三輪機動車（機車，含三輪車）是 `CP_39`（gsrrs_2023_motorcycles()）。
"""
from __future__ import annotations

from dataclasses import dataclass

import openpyxl

from common import download
from sources import who
from sources.countries import to_iso3

GSRRS_URL = (
    "https://cdn.who.int/media/docs/default-source/documents/health-topics/road-traffic-injuries/"
    "gsrrs23-indicators-for-participating-countries-or-territories.xlsx?sfvrsn=88035adb_3"
)

# 道路使用者的分類（WHO 的五類），跟 RS_246 的 Dim1、GSRRS 2023 的 CP_18a–e 對應
USERS = ["car", "moto", "pedestrian", "cyclist", "other"]
RS_246_USERS = {
    "ROADUSERTYPE_RS-DDC-4WHEELS": "car",
    "ROADUSERTYPE_RS-DDC-2OR3WHEELS": "moto",
    "ROADUSERTYPE_RS-DDC-PEDESTRIANS": "pedestrian",
    "ROADUSERTYPE_RS-DDC-CYCLISTS": "cyclist",
    "ROADUSERTYPE_RS-DDC-OTHER": "other",
}
GSRRS_USERS = {"CP_18a": "car", "CP_18b": "moto", "CP_18c": "pedestrian", "CP_18d": "cyclist", "CP_18e": "other"}


def death_rate() -> tuple[dict[str, dict[int, float]], dict[int, float]]:
    """WHO 估計的道路交通死亡率（每 10 萬人）：(各國 {iso3: {年: 值}}, 世界 {年: 值})"""
    return who.gho("RS_198", {"Dim1": "SEX_BTSX"})


@dataclass
class Survey:
    """一個國家在一次調查裡通報的資料"""

    year: int | None  # 死亡人數依道路使用者分的資料年份
    shares: dict[str, float]  # 各類道路使用者占死亡的百分比
    vehicles: float | None = None  # 登記的機動車輛數
    vehicles_year: int | None = None


def _num(v: object) -> float | None:
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def gsrrs_2016() -> dict[str, Survey]:
    """GSRRS 2018 那一輪（GHO 的 RS_246 與 RS_194，CC BY 4.0）"""
    out: dict[str, Survey] = {}
    rows = [r for r in who.gho_rows("RS_246") if r["NumericValue"] is not None]
    # 少數國家有兩個年份（史瓦帝尼 2013、2016），只用最新的一年，不同年份的占比不混在一起
    latest: dict[str, int] = {}
    for row in rows:
        if row["SpatialDimType"] == "COUNTRY" and (iso3 := to_iso3(row["SpatialDim"])):
            latest[iso3] = max(latest.get(iso3, 0), row["TimeDim"])
    for row in rows:
        iso3 = to_iso3(row["SpatialDim"]) if row["SpatialDimType"] == "COUNTRY" else None
        user = RS_246_USERS.get(row["Dim1"] or "")
        if not iso3 or not user or row["TimeDim"] != latest[iso3]:
            continue
        s = out.setdefault(iso3, Survey(row["TimeDim"], {}))
        s.shares[user] = float(row["NumericValue"])
    for row in who.gho_rows("RS_194"):
        iso3 = to_iso3(row["SpatialDim"]) if row["SpatialDimType"] == "COUNTRY" else None
        if not iso3 or row["NumericValue"] is None:
            continue
        s = out.setdefault(iso3, Survey(None, {}))
        s.vehicles, s.vehicles_year = float(row["NumericValue"]), row["TimeDim"]
    return out


def gsrrs_2023() -> dict[str, Survey]:
    """GSRRS 2023 的指標檔（CC BY-NC-SA 3.0 IGO）：第一列是欄位代碼、第二列是欄名，第三列起每國一列"""
    path = download(GSRRS_URL, "who-road", "gsrrs23-indicators.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["Indicators"].iter_rows(values_only=True))
    codes = [str(c) if c is not None else "" for c in rows[0]]
    col = {c: i for i, c in enumerate(codes) if c}
    out: dict[str, Survey] = {}
    for row in rows[2:]:
        iso3 = to_iso3(str(row[col["CP_0"]] or ""))
        if not iso3:
            continue
        year = _num(row[col["CP_16"]])
        shares = {user: v for code, user in GSRRS_USERS.items() if (v := _num(row[col[code]])) is not None}
        vehicles = _num(row[col["CP_36a"]])
        vehicles_year = _num(row[col["CP_118"]])
        out[iso3] = Survey(
            int(year) if year else None,
            shares,
            vehicles if vehicles else None,
            int(vehicles_year) if vehicles_year else None,
        )
    return out


def gsrrs_2023_motorcycles() -> dict[str, tuple[int, float, float]]:
    """GSRRS 2023 的登記機車數（CP_39「Powered 2- and 3-wheelers」）：{iso3: (年份, 機車數, 登記車輛總數)}，
    三個都有值的國家才收（印度沒有填機車數）"""
    path = download(GSRRS_URL, "who-road", "gsrrs23-indicators.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["Indicators"].iter_rows(values_only=True))
    col = {str(c): i for i, c in enumerate(rows[0]) if c is not None}
    out: dict[str, tuple[int, float, float]] = {}
    for row in rows[2:]:
        iso3 = to_iso3(str(row[col["CP_0"]] or ""))
        moto, total, year = (_num(row[col[c]]) for c in ("CP_39", "CP_36a", "CP_118"))
        if iso3 and moto is not None and total and year:
            out[iso3] = (int(year), moto, total)
    return out


def gsrrs_2023_cars() -> dict[str, tuple[int, float, float | None]]:
    """GSRRS 2023 的登記汽車數（CP_38「Cars and 4-wheeled light vehicles」，含皮卡與輕型貨車）：
    {iso3: (年份, 汽車數, 登記車輛總數或 None)}。奧地利沒有填年份（CP_118），當成跟其他國家一樣的 2021 年。"""
    path = download(GSRRS_URL, "who-road", "gsrrs23-indicators.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["Indicators"].iter_rows(values_only=True))
    col = {str(c): i for i, c in enumerate(rows[0]) if c is not None}
    out: dict[str, tuple[int, float, float | None]] = {}
    for row in rows[2:]:
        iso3 = to_iso3(str(row[col["CP_0"]] or ""))
        cars, total, year = (_num(row[col[c]]) for c in ("CP_38", "CP_36a", "CP_118"))
        if iso3 and cars is not None:
            out[iso3] = (int(year or 2021), cars, total)
    return out
