"""SIPRI Military Expenditure Database（https://www.sipri.org/databases/milex）。

不是開放授權：非商業使用、註明出處，轉載少於整個資料集 10% 才算合理使用，超過要向 SIPRI 申請授權
（https://www.sipri.org/about/terms-and-conditions）。整個資料集約 5.5 萬個數值，
所以這裡只取呼叫端指定的年份範圍。含台灣。

Excel 每張表的格式：第 6 列是欄名（Country、Notes、1949…），之後每列一個國家，
區域標題列（Africa、Asia & Oceania…）沒有數值；缺值寫成「...」，國家當年不存在寫成「xxx」。
"""
from __future__ import annotations

import openpyxl

from common import download
from sources.countries import from_name

URL = "https://www.sipri.org/sites/default/files/SIPRI-Milex-data-1949-2025_v1.2.xlsx"

# SIPRI 的國名寫法，pycountry 查不到或會查錯的
NAMES = {
    "Congo, DR": "COD", "Congo, Republic": "COG", "Gambia, The": "GMB", "Korea, North": "PRK", "Korea, South": "KOR",
    "Bosnia and Herzegovina": "BIH", "Central African Republic": "CAF", "Dominican Republic": "DOM",
    "Trinidad and Tobago": "TTO", "Taiwan": "TWN", "Kosovo": "XKX", "Türkiye": "TUR", "Turkey": "TUR",
    "Viet Nam": "VNM", "Laos": "LAO", "Brunei": "BRN", "Czechia": "CZE", "Czech Republic": "CZE",
    "North Macedonia": "MKD", "Eswatini": "SWZ", "Timor Leste": "TLS", "Timor-Leste": "TLS",
    "UAE": "ARE", "United Arab Emirates": "ARE", "USSR": "", "Yemen, North": "", "German DR": "",
    "Czechoslovakia": "", "Yugoslavia": "", "Libya": "LBY", "Iran": "IRN", "Syria": "SYR",
    "Moldova": "MDA", "Russia": "RUS", "Tanzania": "TZA", "Venezuela": "VEN", "Bolivia": "BOL",
    "Cape Verde": "CPV", "Palestine": "PSE", "European Union": "",
}


def _iso3(name: str) -> str | None:
    """ISO alpha-3；刻意略過的（歐盟、已不存在的國家）回傳空字串，查不到回傳 None。"""
    name = name.strip().rstrip("*†‡§¶ ").strip()
    if name in NAMES:
        return NAMES[name]
    return from_name(name)


def table(sheet: str, first_year: int, last_year: int) -> dict[str, dict[int, float]]:
    """回傳 {iso3: {年: 值}}。sheet 例如 'Current US$'（百萬美元）、'Share of GDP'（比例，0.025 = 2.5%）。"""
    path = download(URL, "sipri", URL.rsplit("/", 1)[1])
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = wb[sheet].iter_rows(values_only=True)
    header: list | None = None
    out: dict[str, dict[int, float]] = {}
    unknown = []
    for row in rows:
        if header is None:
            if row[0] == "Country":
                header = list(row)
            continue
        if not row[0] or not any(isinstance(c, (int, float)) for c in row[2:]):
            continue  # 區域標題、空白列
        values = {
            year: float(cell)
            for year, cell in zip(header[2:], row[2:])
            if isinstance(year, int) and first_year <= year <= last_year and isinstance(cell, (int, float))
        }
        if not values:
            continue
        iso3 = _iso3(str(row[0]))
        if iso3:
            out[iso3] = values
        elif iso3 is None:
            unknown.append(str(row[0]))
    if unknown:
        print(f"  SIPRI 無法對應 ISO 的國名（略過）：{', '.join(unknown)}")
    return out


def dataset_size() -> int:
    """整個資料集的數值筆數（各國各表的數值格子，不含封面與註腳），用來確認轉載的比例。"""
    path = download(URL, "sipri", URL.rsplit("/", 1)[1])
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return sum(
        isinstance(cell, (int, float))
        for ws in wb.worksheets
        if ws.title not in ("Front page", "Footnotes")
        for row in ws.iter_rows(values_only=True)
        for cell in row
    )
