"""衛生福利部統計處（政府資料開放授權條款）：補上世界銀行與 WHO 沒有的台灣醫療與疾病資料。

- 國民醫療保健支出（NHE）統計表，表1：經常性醫療保健支出（CHE）占 GDP、平均每人 CHE（元、美元）、
  公部門（各級政府＋全民健保）占 CHE 比重、家庭自付占 CHE 比重。
  CHE = NHE 扣掉資本形成，跟 WHO／OECD 的 current health expenditure 同一個概念。
  每年新版會修訂整個序列，下載網址的 dl 編號也會變：到 https://dep.mohw.gov.tw/DOS/lp-5071-113.html
  找「NNN年國民醫療保健支出統計表」的 .xlsx。
- 衛生福利統計指標（https://dep.mohw.gov.tw/DOS/lp-5083-113.html），原地更新，
  都只有 2000、2005、2010 年與 2015 年以後每年：
  - 表4-2 病床數：平均每萬人口病床數（醫院，不含診所）。
  - 表4-3 醫事人力：每萬人口執業西醫師數（不含中醫師、牙醫師）、執業護理人員數（護理師＋護士）。
  - 表5 法定傳染病：結核病確定病例（新案）。
  - 表7 死因統計：惡性腫瘤的標準化死亡率（WHO 2000 年世界標準人口）。
"""
from __future__ import annotations

import re

import openpyxl

from common import download

NHE_URL = "https://www.mohw.gov.tw/dl-99630-35f4025b-4521-4ff4-84e8-22979a53a7c3.html"  # 113 年版（資料到 2024 年）
BEDS_URL = "https://www.mohw.gov.tw/dl-63138-6c554eda-c839-4443-b334-e8e7c4321b25.html"
WORKFORCE_URL = "https://www.mohw.gov.tw/dl-63141-3fdaa1b6-5654-4605-8abb-f6c00614d0a4.html"
NOTIFIABLE_URL = "https://www.mohw.gov.tw/dl-63144-9e3c90eb-a7d4-4391-b20a-61c84d006d5a.html"
DEATHS_URL = "https://www.mohw.gov.tw/dl-63150-a5f26a9e-9059-4f3c-ab48-f597d4b2b600.html"

# 國民醫療保健支出表1 的欄位（0 起算）
NHE_COLUMNS = {
    "population": 1,  # 年中人口數
    "che-gdp": 9,  # CHE 占 GDP（%）
    "che-per-capita-ntd": 12,  # 平均每人 CHE（元）
    "che-per-capita-usd": 13,  # 平均每人 CHE（美元）
    "public-share": 15,  # 公部門 CHE 占 CHE（%）
    "oop-share": 17,  # 家庭自付占 CHE（%）
}

TAIWAN_NOTE = "台灣不在世界銀行（WHO）資料庫中，改用衛生福利部的資料"
TAIWAN_NOTE_EN = "Taiwan is not in the World Bank (WHO) database; Ministry of Health and Welfare data used instead"
NHE_NOTE = "國民醫療保健支出統計"
NHE_NOTE_EN = "the National Health Expenditure statistics"
NOTES = {
    "health-spending-per-capita": f"{TAIWAN_NOTE}：{NHE_NOTE}的平均每人經常性醫療保健支出（美元）",
    "health-spending-gdp": f"{TAIWAN_NOTE}：{NHE_NOTE}的經常性醫療保健支出占 GDP（不含資本形成，2024 年含資本形成是 7.34%）",
    "oop-share": (
        f"{TAIWAN_NOTE}：{NHE_NOTE}的家庭自付占經常性醫療保健支出比重。"
        "台灣的家庭自付含健保部分負擔、自費醫療、非處方藥，也含營養保健食品與民俗療法，範圍比多數國家寬，可能偏高"
    ),
    "oop-per-capita": (
        f"{TAIWAN_NOTE}：{NHE_NOTE}的平均每人經常性醫療保健支出（元）× 家庭自付比重，"
        "再用 IMF 的購買力平價換算成國際元；家庭自付含營養保健食品與民俗療法，可能偏高"
    ),
    "government-share": f"{TAIWAN_NOTE}：{NHE_NOTE}的公部門（各級政府＋全民健保）占經常性醫療保健支出比重",
    "physicians": f"{TAIWAN_NOTE}：每萬人口執業西醫師數除以 10（不含中醫師、牙醫師），2001–2014 年只有 2005、2010 年",
    "nurses": (
        f"{TAIWAN_NOTE}：每萬人口執業護理人員數（護理師＋護士）除以 10，不含人數很少的助產人員，"
        "2001–2014 年只有 2005、2010 年"
    ),
    "hospital-beds": (
        f"{TAIWAN_NOTE}：醫院的每萬人口病床數除以 10（不含診所的病床），"
        "含急診觀察床、血液透析床等特殊病床，可能比其他國家的算法多；2001–2014 年只有 2005、2010 年"
    ),
}
NOTES_EN = {
    "health-spending-per-capita": (
        f"{TAIWAN_NOTE_EN}: current health expenditure per person from {NHE_NOTE_EN} (US$)"
    ),
    "health-spending-gdp": (
        f"{TAIWAN_NOTE_EN}: current health expenditure as a share of GDP from {NHE_NOTE_EN} "
        "(excluding capital formation; 7.34% in 2024 including it)"
    ),
    "oop-share": (
        f"{TAIWAN_NOTE_EN}: household out-of-pocket payments as a share of current health expenditure from {NHE_NOTE_EN}. "
        "Taiwan counts co-payments, private care and over-the-counter medicines, but also health supplements and "
        "folk remedies, a wider scope than most countries, so the share may be overstated"
    ),
    "oop-per-capita": (
        f"{TAIWAN_NOTE_EN}: current health expenditure per person (NT$) from {NHE_NOTE_EN} × the out-of-pocket share, "
        "converted to international dollars with the IMF's PPP rate; out-of-pocket includes health supplements and "
        "folk remedies, so it may be overstated"
    ),
    "government-share": (
        f"{TAIWAN_NOTE_EN}: the public sector (all levels of government plus National Health Insurance) "
        f"as a share of current health expenditure from {NHE_NOTE_EN}"
    ),
    "physicians": (
        f"{TAIWAN_NOTE_EN}: practising Western-medicine physicians per 10,000 people divided by 10 "
        "(excluding Chinese-medicine practitioners and dentists); for 2001–2014 only 2005 and 2010 are available"
    ),
    "nurses": (
        f"{TAIWAN_NOTE_EN}: practising nurses (registered nurses and nurses) per 10,000 people divided by 10, "
        "excluding the very few midwives; for 2001–2014 only 2005 and 2010 are available"
    ),
    "hospital-beds": (
        f"{TAIWAN_NOTE_EN}: hospital beds per 10,000 people divided by 10 (excluding beds in clinics), including "
        "special beds such as emergency observation and haemodialysis beds, so possibly more than other countries "
        "count; for 2001–2014 only 2005 and 2010 are available"
    ),
}


def _number(cell: object) -> float | None:
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell).replace(",", "").strip())
    except ValueError:
        return None  # 「...」「-」「*」


def _rows(url: str, filename: str, sheet: str | None = None) -> list[tuple]:
    path = download(url, "mohw", filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(wb[sheet or wb.sheetnames[0]].iter_rows(values_only=True))


def _roc_year(cell: object) -> int | None:
    """「113年」→ 2024"""
    text = str(cell or "").strip()
    return int(text[:-1]) + 1911 if re.fullmatch(r"\d+年", text) else None


def _column_series(rows: list[tuple], column: int, scale: float = 1.0) -> dict[int, float]:
    out: dict[int, float] = {}
    for row in rows:
        if (year := _roc_year(row[0])) and (v := _number(row[column])) is not None:
            out[year] = v * scale
    return out


def nhe() -> dict[str, dict[int, float]]:
    """國民醫療保健支出表1：{NHE_COLUMNS 的鍵: {西元年: 值}}。"""
    rows = _rows(NHE_URL, "nhe.xlsx", "表1")
    out = {key: _column_series(rows, column) for key, column in NHE_COLUMNS.items()}
    if not out["che-per-capita-usd"]:
        raise RuntimeError("國民醫療保健支出表1 讀不到資料，表格格式可能變了")
    return out


def _per_10000_column(rows: list[tuple], label: str) -> int:
    """醫事人力表：同一列的 label 出現兩次，第一次是人數，第二次是每萬人口"""
    header = next(r for r in rows if sum(c == label for c in r) == 2)
    return [i for i, c in enumerate(header) if c == label][1]


def _workforce_per_1000(label: str) -> dict[int, float]:
    rows = _rows(WORKFORCE_URL, "physicians.xlsx")
    return _column_series(rows, _per_10000_column(rows, label), 0.1)


def hospital_beds_per_1000() -> dict[int, float]:
    rows = _rows(BEDS_URL, "beds.xlsx")
    # 表頭：第一列「平均每萬人口病床數」，下一列從那一欄開始的第一個「醫院」
    top = next(i for i, r in enumerate(rows) if "平均每萬人口病床數" in r)
    start = rows[top].index("平均每萬人口病床數")
    column = next(i for i, c in enumerate(rows[top + 1]) if i >= start and c == "醫院")
    return _column_series(rows, column, 0.1)


def tuberculosis_cases() -> dict[int, float]:
    """結核病確定病例數（新案，不含再治）。"""
    rows = _rows(NOTIFIABLE_URL, "notifiable.xlsx")
    header = next(r for r in rows if "結核病" in r)
    return _column_series(rows, header.index("結核病"))


def cancer_death_rate() -> dict[int, float]:
    """惡性腫瘤的每十萬人口標準化死亡率（WHO 2000 年世界標準人口）。"""
    rows = _rows(DEATHS_URL, "deaths.xlsx")
    header = next(r for r in rows if "惡性腫瘤" in r)
    return _column_series(rows, header.index("惡性腫瘤"))


def health_series() -> dict[str, dict[int, float]]:
    """健康主題用的台灣序列，鍵是指標檔名（不含要另外換算的人均自付，見 health.py）。"""
    table = nhe()
    return {
        "health-spending-per-capita": table["che-per-capita-usd"],
        "health-spending-gdp": table["che-gdp"],
        "oop-share": table["oop-share"],
        "government-share": table["public-share"],
        "physicians": _workforce_per_1000("西醫師"),
        "nurses": _workforce_per_1000("護理人員"),
        "hospital-beds": hospital_beds_per_1000(),
    }
