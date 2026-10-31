"""研發投入最多的公司（2026-10-31，製作手冊第十批 173）：歐盟執委會 JRC「EU Industrial R&D Investment Scoreboard」的
全球研發投入前 2,000 大公司。使用者決定 736（大型科技公司的總部）、737 合成一個獨立的主題。

- 資料：knowledge-for-policy.ec.europa.eu 的 World2000(2004-2025).zip（CC BY 4.0，執委會的 Legal notice 與報告
  JRC144638 的版權頁），每一版一個 xlsx；用 2013–2025 年版（2012–2024 會計年度），各版收錄家數不同（2,000 或 2,500），
  時間軸一律用前 2,000 大（使用者決定）。研發算在總部所在國（使用者決定，說明寫明），金額是百萬歐元。
- 總部的點（最新一版）：研發前 300 大與台灣的公司用名稱的幾種寫法在 WDQS 查 Wikidata 的總部（P159）座標，國家要跟
  Scoreboard 一樣才收；對不到的前 50 大與台灣公司另外寫 Wikidata 的名稱（MANUAL）。還是對不到的沒有點，左邊的清單
  點了移到那個國家的首都（寫明）。
- 不用英文維基的科技公司營收表：整份來自 Fortune Global 500，Fortune 的條款不准轉載、改作。

輸出（public/data/topics/rd-companies/）：companies.json、rd.json、rd-share.json（各國在前 2,000 大的家數、研發合計、
占前 2,000 大的比例，2012–2024 年）、hq.json（最新一版的 2,000 家，有總部座標的才畫點）。
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import defaultdict

import openpyxl

from common import OUT_DIR, download, indicator_file, now_iso, points_file, write_json
from sources.countries import from_name
from sources.wikidata_profiles import _rows, _sparql

OUT = OUT_DIR / "topics" / "rd-companies"
URL = "https://knowledge-for-policy.ec.europa.eu/sites/default/files/World2000%282004-2025%29.zip"
EDITIONS = range(2013, 2026)  # 2013 年版（2012 會計年度）起是世界前 2,000 或 2,500 大
TOP = 2000
COUNTRY_ALIASES = {"US": "USA", "USA": "USA", "UK": "GBR", "South Korea": "KOR", "Korea": "KOR", "Taiwan": "TWN",
                   "Russia": "RUS", "Cayman Islands": "CYM", "Bermuda": "BMU", "Virgin Islands (British)": "VGB",
                   "Virgin Islands (UK)": "VGB", "Hong Kong": "HKG", "Czech Republic": "CZE", "Turkey": "TUR",
                   "Netherlands": "NLD", "The Netherlands": "NLD", "Taiwan, China": "TWN"}
# 產業（ICB3 short）→ 點的顏色代碼（前端 src/topics/rd-companies.ts 的 SECTORS）
SECTORS = {"ICT producers": 1, "ICT services": 2, "Health industries": 3, "Automobiles & Parts": 4, "Industrials": 5,
           "Aerospace & Defence": 5, "Chemicals": 6, "Energy": 6, "Construction & Materials": 6, "Financial": 6, "Others": 6}
# 名稱的寫法對不到的公司：2025 年版的排名 → (Wikidata 的英文名稱, 允許的國家；None 是不檢查)
MANUAL = {
    6: ("Huawei", "CHN"), 11: ("Merck & Co.", "USA"), 14: ("AstraZeneca", "GBR"), 17: ("Mercedes-Benz Group", "DEU"),
    19: ("Bristol Myers Squibb", "USA"), 21: ("BMW", "DEU"), 27: ("AbbVie", "USA"), 28: ("Robert Bosch GmbH", "DEU"),
    35: ("Stellantis", None),  # 總部在阿姆斯特丹，Scoreboard 算法國
    43: ("Boehringer Ingelheim", "DEU"), 49: ("Nippon Telegraph and Telephone", "JPN"),
    59: ("MediaTek", "TWN"), 363: ("Asus", "TWN"), 864: ("Nanya Technology", "TWN"), 884: ("Lite-On", "TWN"),
    1182: ("Micro-Star International", "TWN"), 1339: ("E Ink Holdings", "TWN"), 1380: ("Gigabyte Technology", "TWN"),
    1802: ("Powertech Technology", "TWN"), 1961: ("GlobalWafers", "TWN"),
}
LEGAL = {"INC", "INCORPORATED", "CORPORATION", "CORP", "CO", "LTD", "LIMITED", "PLC", "AG", "SE", "SA", "S.A", "N.V", "NV",
         "HOLDINGS", "HOLDING", "GROUP", "COMPANY", "K.K", "KK", "A/S", "AB", "ASA", "OYJ", "S.P.A", "SPA", "GMBH", "THE",
         "IND", "INDUSTRY", "INDUSTRIES", "TECHNOLOGY", "TECHNOLOGIES"}
ACRONYMS = {"IBM", "SAP", "BYD", "ABB", "BMW", "NEC", "TDK", "LG", "SK", "AMD", "ARM", "BASF", "GE", "HP", "ZTE", "NXP", "ASML", "KLA"}


def _iso(name: str) -> str | None:
    name = (name or "").strip()
    return COUNTRY_ALIASES.get(name) or from_name(name)


def _edition(z: zipfile.ZipFile, year: int) -> list[dict]:
    """一版 Scoreboard：[{rank, name, iso3, rd, sales, sector}]，只取前 2,000 大"""
    member = next(n for n in z.namelist() if re.search(rf"SB{year}_World\d+\.xlsx$", n))
    ws = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    is_rank = lambda c: isinstance(c, str) and c.strip().lower() in ("rank", "world rank")  # noqa: E731
    head = next(i for i, r in enumerate(rows[:10]) if r and (is_rank(r[0]) or len(r) > 1 and is_rank(r[1])))
    prev = rows[head - 1] if head else ()
    if sum(1 for c in prev if isinstance(c, str) and c.strip()) < 3:  # 上一列是表的標題，不是第一層表頭
        prev = ()
    cols = [f"{(prev[i] if i < len(prev) and prev[i] else '')} {(c or '')}".strip() for i, c in enumerate(rows[head])]
    find = lambda pred: next(i for i, c in enumerate(cols) if pred(c))  # noqa: E731
    i_rank = find(lambda c: "rank" in c.lower())
    i_name = find(lambda c: c.lower() in ("name", "company") or c.lower().endswith(" name") or c.lower().endswith(" company"))
    i_iso = next((i for i, c in enumerate(cols) if c.lower().startswith("iso3")), None)
    i_country = find(lambda c: c.lower().endswith("country"))
    i_rd = find(lambda c: "r&d" in c.lower() and not re.search(r"growth|intensity|ratio|capex", c, re.I))
    i_sales = next((i for i, c in enumerate(cols) if re.search(r"(net )?sales", c, re.I) and not re.search(r"growth", c, re.I)), None)
    i_sector = next((i for i, c in enumerate(cols) if c.lower() == "icb3 short"), None)
    out = []
    for r in rows[head + 1:]:
        if not r or r[i_rank] in (None, "") or not r[i_name]:
            continue
        try:
            rank = int(float(r[i_rank]))
            rd = float(r[i_rd])
        except (TypeError, ValueError):
            continue
        if rank > TOP:
            continue
        iso = (r[i_iso] if i_iso is not None else None) or _iso(str(r[i_country]))
        if not iso:
            raise RuntimeError(f"Scoreboard {year}：對不到國家「{r[i_country]}」")
        sales = r[i_sales] if i_sales is not None else None
        out.append({"rank": rank, "name": str(r[i_name]).strip(), "iso3": iso, "rd": rd,
                    "sales": float(sales) if isinstance(sales, (int, float)) else None,
                    "sector": r[i_sector] if i_sector is not None else None})
    if len(out) < 1990:
        raise RuntimeError(f"Scoreboard {year}：只讀到 {len(out)} 家，版面可能變了")
    return out


def _variants(name: str) -> set[str]:
    words = [w for w in re.sub(r"\s*\(.*?\)", "", name).replace(",", " ").replace("&", " & ").split() if w]

    def strip(ws: list[str]) -> list[str]:
        ws = list(ws)
        while ws and ws[-1].upper().strip(".") in LEGAL:
            ws.pop()
        return ws

    core = strip(words)
    out: set[str] = set()
    for seq in (core, strip(core[:-1]) if len(core) > 1 else core):
        if not seq:
            continue
        t = " ".join(w if w == "&" else w.upper() if w.upper() in ACRONYMS else w.capitalize() for w in seq)
        t = re.sub(r"\.Com$", ".com", t)
        out |= {t, t + " Inc.", t + " Corporation", t.replace(".com", ""), t.replace(" & ", " and ")}
        if len(seq) == 1:
            out.add(seq[0].upper())
    return {v for v in out if len(v) >= 2}


def _hq(names: list[str], prefix: str) -> dict[str, list[tuple[str | None, float, float]]]:
    """Wikidata 的名稱 → [(國家, 經度, 緯度)]：有總部（P159）而總部有座標的項目"""
    values = " ".join(json.dumps(n, ensure_ascii=False) + "@" + lang for n in sorted(set(names)) for lang in ("en", "mul"))
    query = f"""SELECT ?name ?item ?iso (SAMPLE(?c0) AS ?c) WHERE {{
  VALUES ?name {{ {values} }}
  ?item rdfs:label|skos:altLabel ?name ; wdt:P159 ?hq . ?hq wdt:P625 ?c0 .
  OPTIONAL {{ ?item wdt:P17/wdt:P298 ?iso }}
}} GROUP BY ?name ?item ?iso"""
    out: dict[str, list] = defaultdict(list)
    for r in _rows(_sparql(query, f"{prefix}-{hashlib.sha1(values.encode()).hexdigest()[:10]}.json")):
        m = re.fullmatch(r"Point\(([-\d.eE]+) ([-\d.eE]+)\)", r["c"]["value"])
        if m:
            out[r["name"]["value"]].append((r.get("iso", {}).get("value"), float(m.group(1)), float(m.group(2))))
    return out


def _capitals() -> dict[str, tuple[float, float]]:
    """國家小檔案的首都（points_file 格式，經度是差分）"""
    d = json.loads((OUT_DIR / "topics" / "profiles" / "capitals.json").read_text(encoding="utf-8"))
    cols = dict(zip(d["fields"], d["columns"]))
    lons, acc = [], 0
    for v in cols["lon"]:
        acc += v
        lons.append(acc * d["scale"]["lon"])
    isos = [d["strings"]["iso3"][i] if isinstance(d.get("strings"), dict) and "iso3" in d["strings"] else i for i in cols["iso3"]]
    out = {}
    for i, iso in enumerate(isos):
        role = cols["role"][i]
        if iso not in out or role in (0, 1):
            out[iso] = (lons[i], cols["lat"][i] * d["scale"]["lat"])
    return out


def build() -> None:
    z = zipfile.ZipFile(download(URL, "jrc", "World2000-2004-2025.zip"))
    editions = {ed - 1: _edition(z, ed) for ed in EDITIONS}  # 會計年度 = 版次 − 1
    years = sorted(editions)
    count: dict[str, list] = defaultdict(lambda: [0] * len(years))
    rd: dict[str, list] = defaultdict(lambda: [0.0] * len(years))
    totals = []
    for i, y in enumerate(years):
        totals.append(sum(c["rd"] for c in editions[y]))
        for c in editions[y]:
            count[c["iso3"]][i] += 1
            rd[c["iso3"]][i] += c["rd"]
    share = {c: [round(v / t * 100, 3) for v, t in zip(s, totals)] for c, s in rd.items()}
    note = ("研發算在總部登記的國家，不是研發實際發生的地方", "R&D is counted in the country of the company's registered headquarters, not "
            "where the research is done")
    for name, data, unit, digits in (("companies", count, "家", 0), ("rd", rd, "百萬歐元", 1), ("rd-share", share, "%", 3)):
        world = {"companies": [TOP] * len(years), "rd": [round(t, 1) for t in totals], "rd-share": [100.0] * len(years)}[name]
        write_json(OUT / f"{name}.json", indicator_file(
            name, unit, years, {c: [round(v, digits) if digits else int(v) for v in s] for c, s in sorted(data.items())},
            aggregates={"WLD": world}, notes={"WLD": note[0]}, notesEn={"WLD": note[1]}), quiet=True)
    last = years[-1]
    latest = editions[last]
    i = years.index(last)
    top = sorted(count.items(), key=lambda kv: -kv[1][i])[:8]
    print(f"  Scoreboard {years[0]}–{last} 會計年度（前 {TOP} 大）；{last} 年：" + "、".join(
        f"{c} {n[i]} 家 {rd[c][i]:,.0f}" for c, n in top) + f"；台灣 {count['TWN'][i]} 家 {rd['TWN'][i]:,.0f}")

    # 總部的點：前 300 大與台灣的公司
    want = [c for c in latest if c["rank"] <= 300 or c["iso3"] == "TWN"]
    found = _hq([v for c in want for v in _variants(c["name"])], "rd-companies")
    manual = _hq([n for n, _iso in MANUAL.values()], "rd-companies-manual")
    capitals = _capitals()
    rows = []
    matched = 0
    for c in latest:
        hit = None
        if c["rank"] in MANUAL:
            label, iso = MANUAL[c["rank"]]
            hit = next(((lon, lat) for ciso, lon, lat in manual.get(label, []) if iso is None or ciso == iso), None)
        elif c in want:
            hit = next(((lon, lat) for v in sorted(_variants(c["name"])) for ciso, lon, lat in found.get(v, []) if ciso == c["iso3"]), None)
        if hit:
            matched += 1
        lon, lat = hit or capitals.get(c["iso3"], (None, None))
        rows.append({"lon": round(lon, 4) if lon is not None else None, "lat": round(lat, 4) if lat is not None else None,
                     "r": c["rank"], "rd": round(c["rd"], 1), "s": round(c["sales"], 0) if c["sales"] else None,
                     "sec": SECTORS.get(c["sector"] or "Others", 6), "hq": 1 if hit else None, "name": c["name"], "iso3": c["iso3"]})
    fields = ["lon", "lat", "r", "rd", "s", "sec", "hq", "name", "iso3"]
    write_json(OUT / "hq.json", {
        **points_file(rows, fields, {"lon": 0.0001, "lat": 0.0001, "r": 1, "rd": 0.1, "s": 1, "sec": 1, "hq": 1}, strings=("iso3",)),
        "license": "CC BY 4.0",
        "attribution": "European Commission, Joint Research Centre, The 2025 EU Industrial R&D Investment Scoreboard; "
                       "headquarters locations from Wikidata (CC0)",
    })
    top50 = [x for x in rows if x["r"] <= 50]
    tw = [x for x in rows if x["iso3"] == "TWN"]
    print(f"  總部座標：{matched} 家（前 50 大 {sum(1 for x in top50 if x['hq'])}/50、台灣 {sum(1 for x in tw if x['hq'])}/{len(tw)}）")
    write_json(OUT / "manifest.json", {"topic": "rd-companies", "generatedAt": now_iso(), "editions": [years[0] + 1, last + 1]})
