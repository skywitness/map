"""南極的冰下湖（新主題，地理類，2026-10-30）：Wilson 等 2025 用十年的 CryoSat-2 測高找到的南極活動冰下湖。

**資料**：Zenodo 16330565（cc-by-4.0，Wilson_2025_Antarctica_SGL_data_and_code.zip，559,937 bytes，不必登入），
Wilson, S. F. et al. 2025, Detection of 85 new active subglacial lakes in Antarctica from a decade of CryoSat-2 data,
Nature Communications, https://doi.org/10.1038/s41467-025-63773-9 。
- Wilson_2025_Supplement_pt2.xlsx：Supplementary Table 1 是 85 個新的活動湖（一列一湖，每一次灌滿或排空一組
  5 欄：開始、結束（小數年）、冰面高度變化（m）、面積（km²）、體積變化（公式，面積 × 高度 / 1000 km³）），
  Table 2 是 19 個重新分析的既有活動湖（Smith 2009、2017 的名單；一湖好幾列，一列一次）。
- boundaries/：每個湖在不同期間畫的範圍（GeoJSON，WGS84），新的 121 個檔、既有的 35 個檔；每個湖取面積最大的一個。

使用者 2026-10-30 決定：手冊寫的「179 個活動湖」是 Livingstone 2022（授權待確認）的數字，Wilson 的檔只有 104 個，
先做這 104 個；完整的 773 個等授權確認再加。

**輸出**（public/data/topics/subglacial-lakes/）：lakes.json（湖的中心點）、outlines.json（範圍，封閉的環）、
lakes-count.json（空的指標檔）。不到 1 秒。
"""
from __future__ import annotations

import io
import json
import math
import re
import unicodedata
import zipfile

from common import OUT_DIR, download, indicator_file, lines_file, now_iso, points_file, write_json

OUT = OUT_DIR / "topics" / "subglacial-lakes"
URL = "https://zenodo.org/api/records/16330565/files/Wilson_2025_Antarctica_SGL_data_and_code.zip/content"
NEW, EXISTING = 1, 2
# Table 2 的湖名 → 範圍檔的前綴、地圖上的名稱
EXISTING_FILES = {
    "Subglacial Lake Mercer (Mercer 2)": "Mercer2", "Subglacial Lake Conway (Whillans 4)": "Lake_Conway",
    "MacAyeal 1": "MacAyeal1", "MacAyeal 2": "MacAyeal2", "MacAyeal 3": "Mac3",
    "Thwaites 170 (THW170)": "THW170", "Thwaites 142 (THW142)": "THW142", "Thwaites (THW124)": "THW124",
    "Thwaites 70 (THW70)": "THW70", "Haynes Glacier (HG)": "Haynes_Glacier", "Thwaites West (TW)": "Thwaites_West",
    "Cook E2": "CookE2", "Cook W1 (Ninnis 1)": "CookW1", "Cook W2 (Ninnis 2)": "CookW2", "Institute E1": "Institute_E1",
    "KT1 (Kamb Trunk 1)": "Kamb_Trunk1", "L78 (Mercer 1)": "Mercer1", "Rutford 1": "Rutford1", "David 2": "David2",
}


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def ring_area_km2(ring: list[list[float]]) -> float:
    """經緯度的環的面積（等積近似：經度依緯度縮放）"""
    lat0 = sum(p[1] for p in ring) / len(ring)
    k = math.cos(math.radians(lat0))
    a = 0.0
    for (x0, y0), (x1, y1) in zip(ring, ring[1:] + ring[:1]):
        a += x0 * k * y1 - x1 * k * y0
    return abs(a) / 2 * 111.32 ** 2


def cycles_of(cells: list) -> list[tuple[float, float, float, float | None]]:
    """一列（或一組）裡的灌滿／排空：(開始, 結束, 高度變化 m, 面積 km² 或 None)"""
    out = []
    for i in range(0, len(cells) - 3, 5):
        start, end, dh, area = cells[i:i + 4]
        if isinstance(start, (int, float)) and isinstance(end, (int, float)) and isinstance(dh, (int, float)):
            out.append((float(start), float(end), float(dh), float(area) if isinstance(area, (int, float)) else None))
    return out


def build() -> None:
    import openpyxl

    path = download(URL, "wilson2025", "data.zip")
    outer = zipfile.ZipFile(path)
    root = next(n for n in outer.namelist() if n.endswith("Wilson_2025_Supplement_pt2.xlsx") and "__MACOSX" not in n)
    base = root.rsplit("/", 1)[0]
    wb = openpyxl.load_workbook(io.BytesIO(outer.read(root)), read_only=True)

    lakes: list[dict] = []
    t1 = wb["Supplementary Table 1"]
    for row in list(t1.iter_rows(values_only=True))[1:]:
        if not row[0] or not isinstance(row[1], (int, float)):
            continue
        lakes.append({"key": nfc(row[0]), "name": nfc(row[0]).replace("_", " "), "t": NEW, "lat": row[1], "lon": row[2],
                      "sheet": row[3], "thick": row[6], "cycles": cycles_of(list(row[8:]))})
    t2 = wb["Supplementary Table 2"]
    for row in list(t2.iter_rows(values_only=True))[1:]:
        if row[0] and isinstance(row[1], (int, float)):
            name = nfc(row[0]).replace("Subglacial Lake ", "Lake ")
            lakes.append({"key": row[0], "name": name, "t": EXISTING, "lat": row[1], "lon": row[2], "sheet": row[3],
                          "thick": row[5], "cycles": []})
        if lakes and lakes[-1]["t"] == EXISTING and isinstance(row[7], (int, float)):
            lakes[-1]["cycles"] += cycles_of(list(row[7:12]))
    assert sum(l["t"] == NEW for l in lakes) == 85 and sum(l["t"] == EXISTING for l in lakes) == 19, len(lakes)

    # 範圍：每個湖面積最大的一個檔
    outlines: dict[str, tuple[float, list]] = {}
    for zname, prefix_of in (("new_sgl_lake_boundaries.zip", None), ("existing_sgl_lake_boundaries.zip", EXISTING_FILES)):
        inner = zipfile.ZipFile(io.BytesIO(outer.read(f"{base}/boundaries/{zname}")))
        for n in inner.namelist():
            if not n.endswith(".geojson") or "__MACOSX" in n:
                continue
            # zip 沒標 UTF-8 的檔名被當成 cp437 解開（Français 變亂碼），換回來
            try:
                fixed = n.encode("cp437").decode("utf-8")
            except (UnicodeEncodeError, UnicodeDecodeError):
                fixed = n
            stem = nfc(fixed.rsplit("/", 1)[1][:-len(".geojson")])
            if prefix_of is None:
                key = stem.rsplit("_", 2)[0]
            else:
                key = next((k for k, p in prefix_of.items() if re.fullmatch(re.escape(p) + r"_\d{6}_\d{6}_[\d.]+", stem)), None)
                if key is None:
                    continue  # Institute_E1a、E1b 是 E1 的一部分，不另外畫
            rings = []
            for f in json.loads(inner.read(n))["features"]:
                g = f["geometry"]
                polys = g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]
                rings += [poly[0] for poly in polys]
            area = sum(ring_area_km2(r) for r in rings)
            if key not in outlines or area > outlines[key][0]:
                outlines[key] = (area, rings)
    missing = [l["key"] for l in lakes if l["key"] not in outlines]
    print(f"  {len(lakes)} 個湖（新 85、既有 19），有範圍的 {len(lakes) - len(missing)}；沒有範圍：{missing}")

    rows, lines, types, records = [], [], [], []
    for l in sorted(lakes, key=lambda l: l["lon"]):
        area_b = outlines.get(l["key"], (None,))[0]
        areas = [c[3] for c in l["cycles"] if c[3]]
        vols = [abs(c[2]) * c[3] / 1000 for c in l["cycles"] if c[3]]
        cyc = ";".join(f"{c[0]:.1f}:{c[1]:.1f}:{c[2]:g}" for c in l["cycles"])
        row = {"lon": l["lon"], "lat": l["lat"], "name": l["name"], "t": l["t"], "sheet": l["sheet"],
               "h": round(l["thick"]) if isinstance(l["thick"], (int, float)) else None,
               "a": round(area_b or max(areas, default=0) or 0, 1) or None,
               "v": round(max(vols), 3) if vols else None, "c": cyc}
        rows.append(row)
        if l["key"] in outlines:
            for ring in outlines[l["key"]][1]:
                closed = ring if ring[0] == ring[-1] else ring + [ring[0]]
                lines.append([(x, y) for x, y in closed])
                types.append(l["t"])
                records.append({"name": l["name"], "t": l["t"], "ring": 1})
    write_json(OUT / "lakes.json", points_file(rows, ["lon", "lat", "t", "name", "sheet", "h", "a", "v", "c"],
                                               {"lon": 0.001, "lat": 0.001, "t": 1, "h": 1, "a": 0.1, "v": 0.001},
                                               delta=("lon",), strings=("sheet",)))
    write_json(OUT / "outlines.json", lines_file(lines, 0.001, types, records=records,
                                                 record_fields=["name", "t", "ring"], record_scale={"t": 1, "ring": 1}))
    write_json(OUT / "lakes-count.json", indicator_file("lakes-count", "", [2020], {}), quiet=True)
    biggest = sorted(rows, key=lambda r: -(r["a"] or 0))[:5]
    print("  範圍最大：" + "、".join(f"{r['name']} {r['a']} km²" for r in biggest))
    write_json(OUT / "manifest.json", {
        "topic": "subglacial-lakes", "generatedAt": now_iso(),
        "source": "Wilson, S. F. et al. (2025) Detection of 85 new active subglacial lakes in Antarctica from a decade of "
                  "CryoSat-2 data. Nature Communications. Data: Zenodo 16330565 (CC BY 4.0)",
        "lakes": len(rows), "outlines": len(lines),
    })
