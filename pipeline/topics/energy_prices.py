"""電價與油價（2026-10-31，製作手冊第十批 164）：家庭與工業電價、95 無鉛汽油與柴油的零售價（含稅，換成美元），
以及沒選國家時側欄的國際能源價格（世界銀行 Pink Sheet，1960 年起每月）。

使用者決定（2026-10-30、10-31）：
- 歷年，一年一個值；有官方長序列的照做，其他國家補 G20 與台灣周邊（政府或主管機關優先、新聞補），每個值寫出處與日期；
- 換算用當年的匯率（世界銀行 WDI PA.NUS.FCRF，年平均；歐元用歐元區的；台灣用中央銀行的年平均）；
- 台電「各國電價比較」只當台灣的對照，寫在附註。

官方長序列：
- 電價：Eurostat nrg_pc_204（家庭，年用電 2,500–4,999 度，含所有稅）、nrg_pc_205（非家庭，年用電 2,000–19,999 MWh，
  含所有稅），歐元，半年一筆，兩個半年都有才算年平均；美國 EIA《Monthly Energy Review》表 9.8 的住宅、工業平均電價
  （含稅）；台灣：台電「公用售電業售電統計資料」（data.gov.tw 6064）的電燈、電力平均單價。
- 油價：歐盟 Weekly Oil Bulletin 的含稅價（Euro-super 95、柴油，歐元／千公升，非歐元國家已換成歐元），每週的平均；
  美國 EIA《Monthly Energy Review》表 9.4 的高級無鉛（Premium，抗爆指數 91 以上，約等於 RON 95）與公路柴油（含稅，美元／加侖）；
  台灣：能源署油價資訊管理與分析系統「汽柴油參考零售價」的中油 95 無鉛、超級柴油牌價，依天數加權的年平均。
- 其他國家：pipeline/data/energy-prices/research.jsonl（逐筆查的，每筆有出處、日期、是年平均還是某一天的價格）。
"""
from __future__ import annotations

import csv
import json
import statistics
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import common
from common import download, indicator_file, now_iso, write_json
from sources import cbc_fx
from sources.eurostat import URL as EUROSTAT_URL, _iso3 as eurostat_iso3

OUT = Path(__file__).resolve().parent.parent.parent / "public" / "data" / "topics" / "energy-prices"
RESEARCH = Path(__file__).resolve().parent.parent / "data" / "energy-prices" / "research.jsonl"
YEARS = list(range(2005, 2026))
LITRES_PER_GALLON = 3.785411784
WOB_URL = ("https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
           "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
MER = "https://www.eia.gov/totalenergy/data/browser/csv.php?tbl={}"
PINK_SHEET = ("https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
              "CMO-Historical-Data-Monthly.xlsx")
TAIPOWER_SALES = "https://service.taipower.com.tw/data/opendata/apply/file/d003001/001.csv"
TAIPOWER_INTL = "https://service.taipower.com.tw/data/opendata/apply/file/d007009/001.csv"
MOEAEA_RETAIL = "https://www2.moeaea.gov.tw/oil111/Gasoline/RetailPrice/load"
INDICATORS = ("household", "industrial", "gasoline", "diesel")
# 側欄的國際能源價格（Pink Sheet 的欄名 → 代號）
COMMODITIES = {
    "Crude oil, Brent": "brent", "Crude oil, Dubai": "dubai", "Crude oil, WTI": "wti",
    "Natural gas, Europe": "gas-europe", "Natural gas, US": "gas-us", "Liquefied natural gas, Japan": "lng-japan",
    "Coal, Australian": "coal-australia", "Coal, South African **": "coal-south-africa",
}


# ── 匯率 ─────────────────────────────────────────


def exchange_rates() -> tuple[dict[str, dict[int, float]], dict[int, float]]:
    """{iso3: {年: 本國貨幣/美元}}、歐元/美元；台灣用中央銀行的年平均"""
    path = download("https://api.worldbank.org/v2/country/all/indicator/PA.NUS.FCRF?format=json&per_page=20000&date=2000:2025",
                    "worldbank", "PA.NUS.FCRF.json")
    _meta, rows = json.loads(Path(path).read_text(encoding="utf-8"))
    out: dict[str, dict[int, float]] = defaultdict(dict)
    for r in rows:
        if r["value"] is not None and r["countryiso3code"]:
            out[r["countryiso3code"]][int(r["date"])] = float(r["value"])
    out["TWN"] = cbc_fx.ntd_per_usd()
    return out, out["EMU"]


# ── 官方長序列 ─────────────────────────────────────


def eurostat_annual(dataset: str, band: str) -> dict[str, dict[int, float]]:
    """{iso3: {年: 歐元/度}}，兩個半年的平均（只有一個半年的那一年不算）"""
    import gzip

    path = download(EUROSTAT_URL.format(dataset=dataset), "eurostat", f"{dataset}.tsv.gz")
    out: dict[str, dict[int, float]] = {}
    with gzip.open(path, "rt", encoding="utf-8") as f:
        header = f.readline().rstrip("\n").split("\t")
        dims = header[0].split("\\")[0].split(",")
        periods = [h.strip() for h in header[1:]]
        want = {"siec": "E7000", "nrg_cons": band, "unit": "KWH", "tax": "I_TAX", "currency": "EUR"}
        for line in f:
            cells = line.rstrip("\n").split("\t")
            key = dict(zip(dims, cells[0].split(",")))
            if any(key.get(k) != v for k, v in want.items()):
                continue
            iso3 = eurostat_iso3(key["geo"])
            if not iso3:
                continue
            halves: dict[int, list[float]] = defaultdict(list)
            for period, cell in zip(periods, cells[1:]):
                number = cell.strip().split(" ")[0]
                if number not in (":", "") and period.endswith(("-S1", "-S2")):
                    halves[int(period[:4])].append(float(number))
            values = {y: sum(v) / 2 for y, v in halves.items() if len(v) == 2}
            if values:
                out[iso3] = values
    return out


def weekly_oil_bulletin() -> dict[str, dict[str, dict[int, float]]]:
    """{'gasoline'|'diesel': {iso3: {年: 歐元/公升}}}：每週含稅價的平均（一年至少 40 週）"""
    import openpyxl

    path = download(WOB_URL, "eu-oil-bulletin", "Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Prices with taxes"]
    rows = ws.iter_rows(values_only=True)
    names = next(rows)
    columns = {}
    for i, name in enumerate(names):
        if isinstance(name, str) and name.endswith(("_price_with_tax_euro95", "_price_with_tax_diesel")):
            geo = name.split("_", 1)[0]
            if geo in ("EU", "EUR"):
                continue
            iso3 = eurostat_iso3({"GR": "EL"}.get(geo, geo))
            columns[i] = (iso3, "gasoline" if name.endswith("euro95") else "diesel")
    weekly: dict[tuple[str, str], dict[int, list[float]]] = defaultdict(lambda: defaultdict(list))
    for row in rows:
        date = row[0]
        if not hasattr(date, "year"):
            continue
        for i, (iso3, fuel) in columns.items():
            v = row[i]
            if isinstance(v, (int, float)) and v > 0:
                weekly[(iso3, fuel)][date.year].append(v / 1000)
    out: dict[str, dict[str, dict[int, float]]] = {"gasoline": {}, "diesel": {}}
    for (iso3, fuel), years in weekly.items():
        out[fuel][iso3] = {y: statistics.fmean(v) for y, v in years.items() if len(v) >= 40}
    return out


def eia_annual(table: str, msn: str) -> dict[int, float]:
    path = download(MER.format(table), "eia", f"MER-{table}.csv")
    with open(path, encoding="utf-8-sig") as f:
        out = {}
        for r in csv.DictReader(f):
            if r["MSN"] == msn and r["YYYYMM"].endswith("13"):
                try:
                    out[int(r["YYYYMM"][:4])] = float(r["Value"])
                except ValueError:  # Not Available、Not Applicable
                    continue
        return out


def taipower() -> tuple[dict[int, float], dict[int, float]]:
    """台電的電燈、電力平均單價（元/度）；service.taipower.com.tw 從 cloud 時通時斷，下載不到時用 .cache/taipower/ 的檔"""
    path = download(TAIPOWER_SALES, "taipower", "sales-average-price.csv")
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    lighting = {int(r["年度"]): float(r["電燈平均單價(元)"]) for r in rows if r["電燈平均單價(元)"]}
    power = {int(r["年度"]): float(r["電力平均單價(元)"]) for r in rows if r["電力平均單價(元)"]}
    return lighting, power


def taipower_comparison() -> dict[str, tuple[int, int, float, int]]:
    """台電「各國電價比較」：{用電別: (年度, 台灣的排名, 台灣的電價, 國家數)}"""
    path = download(TAIPOWER_INTL, "taipower", "electricity-price-international.csv")
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    out = {}
    for kind in {r["用電別"] for r in rows}:
        sub = [r for r in rows if r["用電別"] == kind]
        tw = next(r for r in sub if r["國家"].replace(" ", "").replace("　", "") in ("台灣", "臺灣", "中華民國"))
        out[kind] = (int(tw["年度"]), int(tw["排名"]), float(tw["平均電價(台幣元/度)"]), len(sub))
    return out


def moeaea_retail() -> dict[str, dict[int, float]]:
    """能源署油價資訊管理與分析系統「汽柴油參考零售價」（網站的政府網站資料開放宣告：無償、可改作，註明出處）：
    中油 95 無鉛（A95）、超級柴油（Achai）每週公告的牌價與施行日，2005-03-10 起。依天數加權成年平均（元/公升），
    只收整年都有牌價的年份（2006 年起）。資料是頁面自己用的 POST（start、end 兩個日期），從 cloud 時通時斷，重試幾次；
    下載過的存在 .cache/moeaea/retail-price.json。回傳 {"gasoline": {年: 元}, "diesel": {年: 元}}
    """
    target = common.CACHE_DIR / "moeaea" / "retail-price.json"
    if not target.exists() or common.REFRESH:
        body = urllib.parse.urlencode({"start": "2005/01/01", "end": date.today().strftime("%Y/%m/%d")}).encode()
        request = urllib.request.Request(MOEAEA_RETAIL, data=body, headers={"User-Agent": common.USER_AGENT})
        print(f"  下載 {MOEAEA_RETAIL}")
        for attempt in range(5):
            try:
                with urllib.request.urlopen(request, timeout=60, context=common._ssl_context(MOEAEA_RETAIL)) as resp:
                    raw = resp.read()
                break
            except (urllib.error.URLError, ConnectionError) as e:
                if attempt == 4:
                    raise
                print(f"  連線失敗（{e}），{5 * (attempt + 1)} 秒後重試")
                time.sleep(5 * (attempt + 1))
        data = json.loads(raw)
        if data.get("res") != "01":
            raise RuntimeError(f"能源署汽柴油參考零售價回傳錯誤：{data.get('msg')}")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
    rows = json.loads(target.read_text(encoding="utf-8"))["data"]["gasoline"]
    out: dict[str, dict[int, float]] = {}
    for indicator, column in (("gasoline", "A95"), ("diesel", "Achai")):
        changes = sorted((date(*map(int, r["Date"].split("/"))), r[column]) for r in rows if r[column] is not None)
        daily: dict[int, list[float]] = defaultdict(list)
        for (start, price), (end, _next) in zip(changes, changes[1:] + [(date.today(), None)]):
            day = start
            while day < end:
                daily[day.year].append(price)
                day += timedelta(days=1)
        out[indicator] = {y: statistics.fmean(v) for y, v in daily.items()
                          if len(v) >= (366 if y % 4 == 0 else 365) and y < date.today().year}
    return out


# ── 逐筆查的其他國家 ─────────────────────────────────


def research() -> list[dict]:
    """逐筆查的價格（2026-10-31 由四個研究代理查、map cloud 整理），一行一筆：
    iso3、indicator、year、value（本國貨幣）、currency、basis（"annual average"，或日期 "YYYY-MM-DD"／"YYYY-MM"）、
    basisKind（日期的意思：on 那一天的價格、since 那一天起施行的、until 施行到那一天的）、
    source_zh、source_en、source_url、source_type（government、regulator、statistics office、utility、news）、quote（原文），
    detail_zh、detail_en（等級、地區、稅這類跟定義不完全一樣的地方，寫進附註），
    有的另有 hint_zh、hint_en（滑過地圖時的短註，例如「不含稅」）；map: false 的只寫附註、地圖不上色
    （使用者 2026-10-31 決定：電價不是全國平均的，城市、單一電力公司、某一級費率），noValue_zh、noValue_en 是地圖留白的原因"""
    if not RESEARCH.exists():
        return []
    return [json.loads(line) for line in RESEARCH.read_text(encoding="utf-8").splitlines() if line.strip()]


# 官方長序列的說明（同一國混用逐筆查的資料時寫進附註）
OFFICIAL = {
    "eurostat": {
        "household": ("Eurostat 年用電 2,500–4,999 度家庭的電價（含所有稅費，上下半年的平均）",
                      "Eurostat price for households using 2,500–4,999 kWh a year (all taxes included, mean of the two half-years)"),
        "industrial": ("Eurostat 年用電 2,000–19,999 MWh 非家庭用戶的電價（含所有稅費，上下半年的平均）",
                       "Eurostat price for non-household consumers using 2,000–19,999 MWh a year (all taxes included, mean of the two half-years)"),
    },
    "wob": {
        "gasoline": ("EU Weekly Oil Bulletin 每週含稅價的年平均", "annual average of weekly prices with taxes from the EU Weekly Oil Bulletin"),
        "diesel": ("EU Weekly Oil Bulletin 每週含稅價的年平均", "annual average of weekly prices with taxes from the EU Weekly Oil Bulletin"),
    },
}


def _years_text(years: list[int], sep: str) -> str:
    """[2015, 2016, 2017, 2019] → 「2015–2017、2019」"""
    runs: list[list[int]] = []
    for y in sorted(years):
        if runs and y == runs[-1][-1] + 1:
            runs[-1].append(y)
        else:
            runs.append([y])
    return sep.join(f"{r[0]}–{r[-1]}" if len(r) > 1 else str(r[0]) for r in runs)


def research_notes(records: list[dict]) -> tuple[list[str], list[str]]:
    """一國一個指標用到的逐筆資料 → 附註：同一個出處的年平均併成一句，某一天的價格逐年寫日期與本國貨幣的價格"""
    zh: list[str] = []
    en: list[str] = []
    annual: dict[tuple, list[int]] = defaultdict(list)
    for r in records:
        if r["basis"] == "annual average":
            annual[(r["source_zh"], r["source_en"], r.get("detail_zh", ""), r.get("detail_en", ""))].append(int(r["year"]))
    for (src_zh, src_en, d_zh, d_en), years in annual.items():
        gap = " " if src_zh[-1].isascii() and src_zh[-1].isalnum() else ""
        zh.append(f"{_years_text(years, '、')} 年：{src_zh}{gap}的年平均{f'，{d_zh}' if d_zh else ''}")
        en.append(f"{_years_text(years, ', ')}: annual average from {src_en}{f'; {d_en}' if d_en else ''}")
    points: dict[tuple, list[dict]] = defaultdict(list)
    for r in sorted((r for r in records if r["basis"] != "annual average"), key=lambda r: int(r["year"])):
        points[(r["source_zh"].split("（")[0] if r["source_type"] == "news" else r["source_zh"], r.get("detail_zh", ""))].append(r)
    when_zh = {"on": "{} 的價格", "since": "{} 起施行的價格", "until": "施行到 {} 的價格"}
    when_en = {"on": "price on {}", "since": "price in force from {}", "until": "price in force until {}"}
    for (_src, d_zh), rs in points.items():
        items_zh = "；".join(f"{r['year']} 年 {r['value']:g} {r['currency']}（{when_zh[r.get('basisKind', 'on')].format(r['basis'])}）"
                            for r in rs)
        items_en = "; ".join(f"{r['year']}: {r['currency']} {r['value']:g} ({when_en[r.get('basisKind', 'on')].format(r['basis'])})"
                             for r in rs)
        sources_zh = "、".join(dict.fromkeys(r["source_zh"] for r in rs))
        sources_en = ", ".join(dict.fromkeys(r["source_en"] for r in rs))
        zh.append(f"{items_zh}（{sources_zh}{f'；{d_zh}' if d_zh else ''}）")
        d_en = rs[0].get("detail_en", "")
        en.append(f"{items_en} ({sources_en}{f'; {d_en}' if d_en else ''})")
    return zh, en


# ── 國際能源價格（側欄） ─────────────────────────────────


def commodity_history() -> dict:
    import openpyxl

    path = download(PINK_SHEET, "worldbank", "CMO-Historical-Data-Monthly.xlsx")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Monthly Prices"]
    rows = list(ws.iter_rows(values_only=True))
    header = next(r for r in rows if r and r[1] == "Crude oil, average")
    cols = {i: COMMODITIES[name] for i, name in enumerate(header) if name in COMMODITIES}
    monthly: dict[str, dict[int, list[float]]] = {k: defaultdict(list) for k in cols.values()}
    last_month = None
    for r in rows:
        if isinstance(r[0], str) and len(r[0]) == 7 and r[0][4] == "M":
            year = int(r[0][:4])
            last_month = r[0]
            for i, key in cols.items():
                if isinstance(r[i], (int, float)):
                    monthly[key][year].append(r[i])
    years = list(range(1960, int(last_month[:4]) + 1))
    # 側欄畫年平均（月資料從 1960 年 1 月起）；最後一年只到 last_month，標出來
    series = {k: [round(statistics.fmean(v[y]), 2) if v.get(y) else None for y in years] for k, v in monthly.items()}
    return {"years": years, "series": series, "lastMonth": f"{last_month[:4]}-{last_month[5:]}"}


# ── 組起來 ─────────────────────────────────────────


def build() -> None:
    rates, eur = exchange_rates()
    usd: dict[str, dict[str, dict[int, float]]] = {k: defaultdict(dict) for k in INDICATORS}
    # 每國每年用了哪個來源（同一國混用官方長序列與逐筆查的，附註寫官方的那幾年）
    origin: dict[str, dict[str, dict[str, list[int]]]] = {k: defaultdict(lambda: defaultdict(list)) for k in INDICATORS}

    def put(indicator: str, iso3: str, year: int, value_usd: float, source: str) -> None:
        if year in YEARS and year not in usd[indicator][iso3]:
            usd[indicator][iso3][year] = value_usd
            origin[indicator][iso3][source].append(year)

    # 歐洲（歐元）
    for indicator, dataset, band in (("household", "nrg_pc_204", "KWH2500-4999"), ("industrial", "nrg_pc_205", "MWH2000-19999")):
        for iso3, s in eurostat_annual(dataset, band).items():
            for y, v in s.items():
                if y in eur:
                    put(indicator, iso3, y, v / eur[y], "eurostat")
    for fuel, by in weekly_oil_bulletin().items():
        for iso3, s in by.items():
            for y, v in s.items():
                if y in eur:
                    put(fuel, iso3, y, v / eur[y], "wob")
    # 美國（美元）
    for indicator, table, msn, scale in (("household", "T09.08", "ESRCUUS", 0.01), ("industrial", "T09.08", "ESICUUS", 0.01),
                                         ("gasoline", "T09.04", "PUUCUUS", 1 / LITRES_PER_GALLON),
                                         ("diesel", "T09.04", "DFONUUS", 1 / LITRES_PER_GALLON)):
        for y, v in eia_annual(table, msn).items():
            put(indicator, "USA", y, v * scale, "eia")
    # 台灣（新臺幣）
    lighting, power = taipower()
    for indicator, s in (("household", lighting), ("industrial", power)):
        for y, v in s.items():
            if y in rates["TWN"]:
                put(indicator, "TWN", y, v / rates["TWN"][y], "taipower")
    tw_fuel = moeaea_retail()
    for indicator, s in tw_fuel.items():
        for y, v in s.items():
            if y in rates["TWN"]:
                put(indicator, "TWN", y, v / rates["TWN"][y], "moeaea")

    # 其他國家（逐筆查的）：同一國同一年有好幾筆時，年平均優先，再來政府的優先
    notes: dict[str, dict[str, list[str]]] = {k: defaultdict(list) for k in INDICATORS}
    notes_en: dict[str, dict[str, list[str]]] = {k: defaultdict(list) for k in INDICATORS}
    rank = {"government": 0, "regulator": 0, "statistics office": 0, "utility": 1, "news": 2}
    records = sorted((r for r in research() if "indicator" in r),
                     key=lambda r: (r["basis"] != "annual average", rank.get(r["source_type"], 3)))
    skipped = []
    used: dict[tuple[str, str], list[dict]] = defaultdict(list)
    hints: dict[str, dict[str, tuple[str, str]]] = {k: {} for k in INDICATORS}
    no_value: dict[str, dict[str, tuple[str, str]]] = {k: {} for k in INDICATORS}
    for r in records:
        iso3, indicator, year = r["iso3"], r["indicator"], int(r["year"])
        if r.get("map") is False:
            used[(iso3, indicator)].append(r)
            no_value[indicator][iso3] = (r["noValue_zh"], r["noValue_en"])
            continue
        if year in usd[indicator].get(iso3, {}):
            continue
        rate = 1.0 if r["currency"] == "USD" else rates.get(iso3, {}).get(year)
        if rate is None:
            skipped.append((iso3, indicator, year, "沒有匯率"))
            continue
        put(indicator, iso3, year, float(r["value"]) / rate, "research")
        used[(iso3, indicator)].append(r)
        if "hint_zh" in r:
            hints[indicator][iso3] = (r["hint_zh"], r["hint_en"])
    for (iso3, indicator), rs in sorted(used.items()):
        zh, en = research_notes(rs)
        for source, years in origin[indicator][iso3].items():
            if source in OFFICIAL:
                zh.insert(0, f"{_years_text(years, '、')} 年：{OFFICIAL[source][indicator][0]}")
                en.insert(0, f"{_years_text(years, ', ')}: {OFFICIAL[source][indicator][1]}")
        notes[indicator][iso3].extend(zh)
        notes_en[indicator][iso3].extend(en)

    # 台灣的附註：台電「各國電價比較」只當對照（使用者 2026-10-31 決定）
    comparison = taipower_comparison()
    for indicator, kind in (("household", "住宅"), ("industrial", "工業")):
        hit = next((v for k, v in comparison.items() if kind in k), None)
        if hit:
            year, place, price, n = hit
            notes[indicator]["TWN"].insert(0, f"地圖用台電的{'電燈' if indicator == 'household' else '電力'}平均單價。台電「各國電價比較」{year} 年"
                                              f"{kind}用電台灣每度 {price:g} 元，{n} 國裡排第 {place}（台電沒有寫各國數字的出處與換算匯率，只當對照）。")
            notes_en[indicator]["TWN"].insert(0, f"The map uses Taipower’s average unit price for {'lighting' if indicator == 'household' else 'power'} customers. "
                                                 f"Taipower’s international comparison for {year} puts Taiwan’s {'residential' if indicator == 'household' else 'industrial'} "
                                                 f"price at NT${price:g} per kWh, No. {place} of {n} (sources and exchange rates not given; for reference only).")

    for indicator, grade, grade_en in (("gasoline", "95 無鉛汽油", "95-octane unleaded"), ("diesel", "超級柴油", "premium diesel")):
        s = tw_fuel[indicator]
        last = max(s)
        notes[indicator]["TWN"].insert(0, f"地圖用中油{' ' if grade[0].isdigit() else ''}{grade}的牌價，依天數加權的年平均（能源署汽柴油參考零售價）；{last} 年 {s[last]:.2f} 元/公升。")
        notes_en[indicator]["TWN"].insert(0, f"The map uses CPC’s posted price for {grade_en}, averaged over the days each price was in force "
                                             f"(Energy Administration reference retail prices); NT${s[last]:.2f} per litre in {last}.")

    OUT.mkdir(parents=True, exist_ok=True)
    for indicator in INDICATORS:
        unit = "美元/度" if indicator in ("household", "industrial") else "美元/公升"
        values = {c: [round(s[y], 3) if y in s else None for y in YEARS] for c, s in sorted(usd[indicator].items())}
        blank = {c: v for c, v in no_value[indicator].items() if c not in values}
        write_json(OUT / f"{indicator}.json", indicator_file(
            indicator, unit, YEARS, values, latestWithin=3,
            notes={c: "；".join(v) for c, v in notes[indicator].items()},
            notesEn={c: "; ".join(v) for c, v in notes_en[indicator].items()},
            hints={c: v[0] for c, v in hints[indicator].items()}, hintsEn={c: v[1] for c, v in hints[indicator].items()},
            noValue={c: v[0] for c, v in blank.items()}, noValueEn={c: v[1] for c, v in blank.items()}))
        last = {c: s[max(s)] for c, s in usd[indicator].items() if s}
        print(f"  {indicator}：{len(values)} 國；最新一年最高 {sorted(((round(v, 3), c) for c, v in last.items()), reverse=True)[:5]}，"
              f"台灣 {round(last.get('TWN', 0), 3)}")
    write_json(OUT / "commodities.json", commodity_history())
    if skipped:
        print(f"  沒有匯率而略過的 {len(skipped)} 筆：{skipped[:8]}")
    write_json(OUT / "manifest.json", {"topic": "energy-prices", "generatedAt": now_iso()})
